"""Bounded, in-memory readers. Uploaded bytes are never written to disk."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 200
MAX_CELLS = 2_000_000
MAX_XLSX_EXPANDED_BYTES = 64 * 1024 * 1024


class DataLoadError(ValueError):
    """A safe, user-facing input error (never includes cell contents)."""


@dataclass(frozen=True)
class LoadedData:
    frame: pd.DataFrame
    filename: str
    size_bytes: int
    encoding: str | None = None
    sheet: str | None = None


def normalize_missing(frame: pd.DataFrame) -> pd.DataFrame:
    """Return a copy: nulls, empty strings and whitespace-only strings are missing."""
    return frame.replace(r"^\s*$", pd.NA, regex=True).copy(deep=True)


def _validate(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise DataLoadError("仅支持 .csv 和 .xlsx 文件，不支持 .xls 或其他格式。")
    if not data:
        raise DataLoadError("文件为空，请选择包含表头和数据的文件。")
    if len(data) > MAX_FILE_BYTES:
        raise DataLoadError("文件超过 10 MiB 上限，请拆分后重试。")
    return suffix


def _check_shape(rows: int, columns: int) -> None:
    if rows > MAX_ROWS or columns > MAX_COLUMNS or rows * columns > MAX_CELLS:
        raise DataLoadError("数据规模过大：最多 100,000 行、200 列、2,000,000 个单元格，请拆分文件。")


def _headers(values: list) -> list[str]:
    headers = [str(v).strip() if v is not None else "" for v in values]
    if not headers or not any(headers):
        raise DataLoadError("未找到有效表头，请确保第一行包含字段名称。")
    if any(not h for h in headers):
        raise DataLoadError("存在空白字段名称，请补全第一行的表头。")
    if len(set(headers)) != len(headers):
        raise DataLoadError("字段名称重复，请为每列设置唯一名称。")
    _check_shape(0, len(headers))
    return headers


def _frame(rows: list, headers: list[str]) -> pd.DataFrame:
    frame = normalize_missing(pd.DataFrame(rows, columns=headers))
    # Infer numeric columns only when every nonmissing value can be converted.
    # Leading-zero strings stay text, preserving common identifiers.
    for column in frame:
        values = frame[column].dropna()
        if values.empty or values.map(lambda x: isinstance(x, (bool,))).any():
            continue
        if values.map(lambda x: isinstance(x, str) and len(x) > 1
                      and x.startswith("0") and x.isdigit()).any():
            continue
        if pd.api.types.is_datetime64_any_dtype(frame[column]):
            continue
        try:
            frame[column] = pd.to_numeric(frame[column], errors="raise")
        except (ValueError, TypeError):
            pass
    return frame


def _csv(data: bytes) -> tuple[pd.DataFrame, str]:
    text = None
    encoding = "utf-8-sig" if data.startswith(b"\xef\xbb\xbf") else "utf-8"
    for candidate in [encoding, "gb18030"]:
        try:
            text = data.decode(candidate, errors="strict")
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise DataLoadError("无法识别 CSV 编码，请另存为 UTF-8、UTF-8-SIG 或 GB18030。")
    if any(ord(c) < 32 and c not in "\t\r\n" for c in text):
        raise DataLoadError("CSV 含有不支持的控制字符或二进制内容，请检查格式与编码。")
    try:
        reader = csv.reader(StringIO(text, newline=""), strict=True)
        headers = _headers(next(reader, []))
        rows = []
        for row in reader:
            # A physical blank line is retained as an entirely missing record.
            if not row or (len(row) == 1 and not row[0].strip()):
                row = [None] * len(headers)
            if len(row) != len(headers):
                raise DataLoadError("CSV 行的字段数与表头不一致，请检查分隔符、引号或损坏行。")
            _check_shape(len(rows) + 1, len(headers))
            rows.append(row)
        return _frame(rows, headers), encoding
    except csv.Error as exc:
        raise DataLoadError("CSV 解析失败，请检查引号、分隔符及字段长度。") from exc


def _workbook(data: bytes):
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_EXPANDED_BYTES:
                raise DataLoadError("XLSX 解压后超过 64 MiB 安全上限，请精简工作簿。")
            if "xl/workbook.xml" not in archive.namelist():
                raise DataLoadError("文件不是有效的 XLSX 工作簿。")
        # Keep formulas as text: cached Excel formula values can be absent/stale.
        return load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except DataLoadError:
        raise
    except Exception as exc:
        raise DataLoadError("XLSX 读取失败：文件可能损坏、加密或格式不正确。") from exc


def list_sheets(data: bytes, filename: str) -> list[str]:
    if _validate(data, filename) != ".xlsx":
        raise DataLoadError("工作表选择仅适用于 XLSX 文件。")
    workbook = _workbook(data)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def load_data(data: bytes, filename: str, sheet: str | None = None) -> LoadedData:
    suffix = _validate(data, filename)
    # Avoid storing a supplied client-side path in a report.
    safe_name = filename.replace("\\", "/").rsplit("/", 1)[-1]
    if suffix == ".csv":
        frame, encoding = _csv(data)
        return LoadedData(frame, safe_name, len(data), encoding=encoding)
    workbook = _workbook(data)
    try:
        selected = sheet if sheet is not None else workbook.sheetnames[0]
        if selected not in workbook.sheetnames:
            raise DataLoadError("所选工作表不存在，请重新选择。")
        worksheet = workbook[selected]
        if worksheet.max_row and worksheet.max_column:
            _check_shape(max(0, worksheet.max_row - 1), worksheet.max_column)
        reader = worksheet.iter_rows(values_only=True)
        headers = _headers(list(next(reader, [])))
        rows = []
        for row in reader:
            _check_shape(len(rows) + 1, len(headers))
            if len(row) != len(headers):
                raise DataLoadError("XLSX 数据列数与表头不一致。")
            rows.append(row)
        return LoadedData(_frame(rows, headers), safe_name, len(data), sheet=selected)
    except DataLoadError:
        raise
    except Exception as exc:
        raise DataLoadError("工作表解析失败，请检查文件内容和格式。") from exc
    finally:
        workbook.close()
