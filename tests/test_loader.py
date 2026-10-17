from io import BytesIO

import pytest
from openpyxl import Workbook

from src.loader import DataLoadError, list_sheets, load_data


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "gb18030"])
def test_csv_chinese_encodings(encoding):
    loaded = load_data("城市,金额\n杭州,12\n成都,24\n".encode(encoding), "table.csv")
    assert loaded.frame.to_dict("list") == {"城市": ["杭州", "成都"], "金额": [12, 24]}
    assert loaded.encoding == encoding


def workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "第一张"
    sheet.append(["城市", "金额"])
    sheet.append(["上海", 30])
    second = workbook.create_sheet("第二张")
    second.append(["值"])
    second.append([42])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_xlsx_sheets():
    data = workbook_bytes()
    assert list_sheets(data, "data.xlsx") == ["第一张", "第二张"]
    assert load_data(data, "data.xlsx").frame.iloc[0, 0] == "上海"
    assert load_data(data, "data.xlsx", "第二张").frame.iloc[0, 0] == 42
    with pytest.raises(DataLoadError, match="不存在"):
        load_data(data, "data.xlsx", "missing")


@pytest.mark.parametrize("data,name", [(b"", "x.csv"), (b"", "x.xlsx"), (b"a\n1", "x.txt"),
    (b"not an excel", "x.xlsx"), (b"a,b\n1,2,3", "x.csv"), (b'a\n"unfinished', "x.csv"),
    (b"a,a\n1,2", "x.csv"), (b"a,\n1,2", "x.csv"), (b"\x00\x00", "x.csv"),
    (b"\xff", "x.csv"), (b" \n", "x.csv")])
def test_bad_files(data, name):
    with pytest.raises(DataLoadError):
        load_data(data, name)


def test_size_and_dimensions(monkeypatch):
    monkeypatch.setattr("src.loader.MAX_FILE_BYTES", 5)
    with pytest.raises(DataLoadError, match="10 MiB"):
        load_data(b"a\n123456", "x.csv")
    monkeypatch.setattr("src.loader.MAX_FILE_BYTES", 100)
    monkeypatch.setattr("src.loader.MAX_ROWS", 1)
    with pytest.raises(DataLoadError, match="规模"):
        load_data(b"a\n1\n2", "x.csv")


def test_blank_records_header_only_and_identifiers():
    frame = load_data(b"id,value\n001,3\n\n , \n002,4", "x.csv").frame
    assert frame.shape == (4, 2)
    assert frame.iloc[1:3].isna().all().all()
    assert frame.iloc[0, 0] == "001"
    assert load_data(b"a,b\n", "x.csv").frame.empty


def test_literal_na_and_quoted_newline():
    frame = load_data(b'a,b\nNA,"two\nlines"\nNULL,ok', "x.csv").frame
    assert frame.iloc[0].tolist() == ["NA", "two\nlines"]
    assert frame.iloc[1, 0] == "NULL"


def test_xlsx_expansion_limit(monkeypatch):
    monkeypatch.setattr("src.loader.MAX_XLSX_EXPANDED_BYTES", 1)
    with pytest.raises(DataLoadError, match="解压"):
        load_data(workbook_bytes(), "x.xlsx")


def test_xlsx_formula_is_text():
    book = Workbook()
    book.active.append(["formula"])
    book.active.append(["=1+1"])
    stream = BytesIO()
    book.save(stream)
    assert load_data(stream.getvalue(), "x.xlsx").frame.iloc[0, 0] == "=1+1"


def test_filename_is_basename():
    assert load_data(b"a\n1", r"C:\private\x.csv").filename == "x.csv"
